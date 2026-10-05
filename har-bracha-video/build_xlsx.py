# -*- coding: utf-8 -*-
import json,os,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from PIL import Image,ImageOps
from catalog_data import C,X,VID
from scenes import S,timeline,BEAT
SCR="/tmp/claude-0/-home-user-solaredge-ops/73fe8d55-5833-5852-9694-53b5adaa6cc2/scratchpad"
files_rel=json.load(open(f"{SCR}/order.json")); files=[f"{SCR}/{p}" for p in files_rel]; titles=json.load(open(f"{SCR}/files.json"))
san=lambda t:t.replace(' ','_').replace('(','').replace(')','')
orig={}
for folder,v in titles.items():
    for fid,t in v: orig[f"src/{folder}/{san(t)}"]=t
used={}
for s in S:
    for i in s[3]:
        if isinstance(i,int): used.setdefault(i,[]).append(s[0])
HDR=PatternFill("solid",fgColor="1F4E79");HF=Font(bold=True,color="FFFFFF",name="Arial",size=11)
thin=Side(style="thin",color="BBBBBB");B=Border(left=thin,right=thin,top=thin,bottom=thin)
AFILL={"A":"C6EFCE","B":"FFEB9C","C":"EDEDED","X":"FFC7CE"}
def sheet(wb,name,headers,widths):
    ws=wb.create_sheet(name); ws.sheet_view.rightToLeft=True
    ws.append(headers)
    for c,w in enumerate(widths,1):
        ws.column_dimensions[get_column_letter(c)].width=w
        cell=ws.cell(1,c); cell.fill=HDR;cell.font=HF;cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True);cell.border=B
    ws.row_dimensions[1].height=32; ws.freeze_panes="B2"; return ws
def fmt(ws,start=2):
    for r in ws.iter_rows(min_row=start):
        for c in r:
            c.alignment=Alignment(vertical="top",wrap_text=True,horizontal="right");c.border=B;c.font=Font(name="Arial",size=10)
wb=Workbook();wb.remove(wb.active)
# ---------- 1 קטלוג
ws=sheet(wb,"1 קטלוג תמונות",["תצוגה","שם קובץ","קטגוריה מזוהה","תת-נושא","ציון רגשי (1-10)","ציון איכות חזותית (1-10)","משך מומלץ (שנ')","עדיפות","בשימוש בסצנות","סוג צילום","כיוון","רזולוציה","הערות","אישור פרסום (הורים) – למילוי"],[16,34,16,22,10,10,10,9,14,18,8,12,40,16])
os.makedirs("thumbs_tmp",exist_ok=True)
rows=[]
for i,f in enumerate(files):
    f_rel=files_rel[i]; im=ImageOps.exif_transpose(Image.open(f)); w,h=im.size
    ori="אנכית" if h>w*1.05 else ("מרובעת" if abs(h-w)<=w*.05 else "אופקית")
    name=orig.get(f_rel,os.path.basename(f))
    if i in C:
        cat,sub,shot,e,q,d,p,note=C[i]
        folder=f_rel.split("/")[1]
        sc_d=next((sc[2] for sc in S if i in sc[3]),None)
        u=", ".join(used.get(i,[])) or "—"
        if sc_d is None: p="C"; note=(note+" | " if note else "")+"לא בשימוש בסרט (כפילות/דמיון לתמונה אחרת או חולשה)"
        rows.append((i,[None,name,folder if cat==folder else f"{folder} ({cat})",sub,e,q,sc_d if sc_d else "—",p,u,shot,ori,f"{w}x{h}",note,""],p,im))
    else:
        cat,note=X[i]
        rows.append((i,[None,name,f_rel.split("/")[1],cat,"—","—","—","לא לשימוש","—","—",ori,f"{w}x{h}",note,""],"X",im))
for v,(t,e,q,d,p,note) in VID.items():
    rows.append((None,[None,os.path.basename(v),v.split('/')[0],t,e,q,d if d else "—",p,"M28" if p=="A" else "—","וידאו","אופקית","848x480",note,""],p,None))
for n,(i,vals,p,im) in enumerate(rows,2):
    ws.append(vals)
    if im is not None:
        im2=im.copy(); im2.thumbnail((110,85)); tp=f"thumbs_tmp/{i}.jpg"; im2.convert("RGB").save(tp,quality=70)
        xi=XLImage(tp); ws.add_image(xi,f"A{n}")
    ws.row_dimensions[n].height=68
    ws.cell(n,8).fill=PatternFill("solid",fgColor=AFILL.get(p,"FFFFFF"))
fmt(ws)
ws.auto_filter.ref=f"A1:N{ws.max_row}"
# ---------- 2 סטוריבורד
ws=sheet(wb,"2 סטוריבורד",["ציר זמן","מס' סצנה","משך (שנ')","תמונות בשימוש","סוג מעבר","תנועת מצלמה","טקסט על המסך","קריינות","סימן מוזיקלי","מערך","סוג צילום","פרק בסיפור"],[14,9,9,22,24,22,32,12,40,14,22,14])
tl,total=timeline()
CAM={"push":"Push-in איטי (100%→112%)","pull":"Pull-out / זום החוצה (114%→100%)","panR":"Ken Burns – פאן ימינה (זום 108%)","panL":"Ken Burns – פאן שמאלה (זום 108%)","hold":"החזקה עם דריפט כמעט בלתי-מורגש (100%→103%)","punch":"Punch-in מהיר (100%→106%) על הפעימה","video":"וידאו חי + push עדין (100%→104%)","dim":"רקע איטי, בהירות 30%","logo":"אנימציית לוגו (ראו גיליון 5)","tiltUp":"Tilt-up"}
def tc(x):return f"{int(x//60)}:{x%60:05.2f}"
for s,(sid,a,b) in zip(S,tl):
    imgs=[]
    for i in s[3]:
        if isinstance(i,int): imgs.append(os.path.basename(files[i]).replace("WhatsApp_Image_","IMG ")+f" [#{i}]")
        else: imgs.append({"V1":"וידאו: ילד מנגן פסנתר","BG1":"רקע גרדיינט כחול (גרפי)","BG2":"רקע גרדיינט סגול (גרפי)","BG3":"רקע גרדיינט ירוק (גרפי)","LOGO_HB":"לוגו הר ברכה (מקור 2.png)","LOGO_MATNAS":"לוגו מתנ\"ס (מקור 1.jpg, ללא שומרון)","LOGO_BOTH":"שני הלוגואים"}[i])
    ws.append([f"{tc(a)}–{tc(b)}",s[0],s[2],"\n".join(imgs),s[4],CAM[s[5]],s[7],"—",s[8],"דיפטיך (2 תמונות)" if len(s[3])==2 else "יחיד",s[6],s[1]])
fmt(ws)
r=ws.max_row+2
ws.cell(r,1,f"סה\"כ אורך: {total:.0f} שנ' | {len(S)} סצנות | 30fps | 1920x1080 | BPM 120 (פעימה=0.5 שנ'; כל חיתוך על פעימה)").font=Font(bold=True)
# ---------- 3 עריכה
ws=sheet(wb,"3 הוראות עריכה",["סצנה","זום","פאן","חיתוך (Crop)","מהירות","פייד","אפקטים חזותיים","נקודת מיקוד"],[8,24,26,34,12,22,38,28])
ZOOM={"push":("100% → 112%","מרכז, איטי"),"pull":("114% → 100%","מרכז, איטי"),"panR":("108% קבוע","שמאל → ימין, 6% מרוחב הפריים"),"panL":("108% קבוע","ימין → שמאל, 6% מרוחב הפריים"),"hold":("100% → 103%","ללא פאן / דריפט אלכסוני 1%"),"punch":("100% → 106% ב-0.5 שנ'","קפיצה קצרה על הפעימה"),"video":("100% → 104%","ללא"),"dim":("100% → 105%, איטי מאד","ללא"),"logo":("ראו גיליון 5","ראו גיליון 5")}
for s in S:
    z,p=ZOOM[s[5]]
    n=len(s[3]); sid=s[0]
    if s[5]=="logo": crop="לוגו על רקע בהיר/לבן, מרווח ביטחון 8%"
    elif sid=="C03": crop="גרדיינט כחול-כהה (#0B2545 → #13315C)"
    elif n==2: crop="דיפטיך: שני פאנלים 944x1080 עם מרווח 8px; חיתוך אנכי-לפנים, שערים נוגדי-פאן (אחד עולה, אחד יורד)"
    elif s[2]<=1.2: crop="16:9 ממורכז סביב נקודת המיקוד; לא לחתוך ראשים"
    else: crop="16:9 + מרווח זום 12%; חיתוך חותמות תאריך (אם יש)"
    spd="100%" if s[5]!="video" else "100% (קול מקורי)"
    fade="פייד שחור 0.8 שנ'" if sid=="S01" else ("עמעום תוך 1.2 שנ' לשחור בסוף" if sid=="L03" else ("הצלבה 0.5 שנ'" if "הצלבה" in s[4] else "חיתוך ישיר"))
    fx="גריידינג חם (+8 חמימות, +5 ניגודיות), ויניטה עדינה"
    if s[5]=="punch": fx+="; רעידה קצרה 2px על ההיט"
    if s[4].startswith("Whip"): fx+="; טשטוש תנועה אופקי 6 פריימים"
    if s[4].startswith("זום"): fx+="; זום-דרך עם טשטוש רדיאלי קל"
    if s[5]=="dim": fx="בהירות 30%, טשטוש רקע 2px, חלקיקי אור (bokeh) איטיים"
    if s[5]=="logo": fx="ראו גיליון 5"
    ws.append([sid,z,p,crop,spd,fade,fx,s[9]])
fmt(ws)
# ---------- 4 ציר זמן
ws=sheet(wb,"4 ציר זמן הפקה",["שנייה","סצנה","פרק","תמונה/אלמנט פעיל","תנועה","טקסט על המסך","הערת מוזיקה","פעימות (0.5 שנ')"],[8,8,14,34,26,34,40,16])
from scenes import start as _start
sec_music=[(0,"אינטרו – ארפג'ו פסנתר, מחיאות, שייקר; קיק נכנס ב-0:08"),(12,"בילד-אפ – קיק על כל פעימה, גלגול סנר, ריזר"),(_start("M08"),"דרופ! קיק+בס+סינת'+הוק, 120 BPM, סיידצ'יין"),(_start("M27"),"צמיחה – פסנתר + פד + קיק בינוני (עדיין קצבי)"),(_start("M34"),"קרשנדו – גלגול סנר + ריזר"),(_start("C01"),"סיום – פסנתר דליל, שקט בין המשפטים, ריזר ב-C03"),(_start("L01"),"דרופ סופי + הוק; היט נוסף בפיצול הלוגו; אקורד אחרון")]
def msec(t):
    m=""
    for a,b in sec_music:
        if t>=a: m=b
    return m
for sec in range(int(total)):
    t0,t1=sec,sec+1
    act=[(s,a,b) for s,(sid,a,b) in zip(S,tl) if a<t1-1e-9 and b>t0+1e-9]
    ids=" / ".join(s[0] for s,_,_ in act)
    el=[]
    for s,_,_ in act:
        el.append(", ".join(f"#{i}" if isinstance(i,int) else str(i) for i in s[3]))
    txt=" | ".join(s[7] for s,a,b in act if s[7] and a<t1 and b>t0 and (a>=t0 or t0<a+4))
    ws.append([f"{tc(sec)}",ids,act[0][0][1],"; ".join(el),"; ".join(CAM[s[5]].split(" (")[0] for s,_,_ in act),txt,msec(sec),f"{sec/BEAT:.0f}–{(sec+1)/BEAT:.0f}"])
fmt(ws)
# ---------- 5 לוגו
ws=sheet(wb,"5 רצף לוגואים",["שלב","זמן (בתוך הרצף)","מה קורה","פירוט טכני"],[8,16,48,70])
LOGO=[
("L-0","0:00–0:03.4","שני הלוגואים מתחילים ביחד באותה סצנה: ילדי המתנ\"ס (תלת-ממדיים, עם פנים, צל ותנועה) קופצים על הגבעה, ומאחוריהם קו הרקיע של הר ברכה בעדינות","זום התחלתי חזק על הדמויות (2.3x) שיורד ל-1x; הדמויות מקבלות הצללה (volumetric), עיניים, חיוך, לחיים, צל על הקרקע, קפיצה כל שנייה (2 פעימות), סקווש וסטרץ', עיוות 'פיתול' במותניים לתנועת זרועות; שאר הלוגו נבנה סביבן: גבעה, עצים, ספר, שמש, קו מקווקו, לב; רקע הר ברכה (בניינים, בית, עצים) בשקיפות 42%"),
("L-1","0:03.4–0:04.9","הפירוד: קרן אור אנכית נפתחת במרכז המסך, והסצנה מתחלקת על המסך לשני לוגואים: מתנ\"ס שמאלה והר ברכה ימינה","כל שכבה זזה בנתיב קשת עם overshoot ורוטציה קלה; שכבות הר ברכה מתמלאות לשקיפות מלאה; גבעה וילדי הר ברכה נכנסים"),
("L-2","0:03.4–0:05.2","הילדים 'מתפשטים' ללוגו: הצללה, פנים וצל נעלמים בהדרגה והופכים לקווים השטוחים של הלוגו","פרמטר flat מ-0 ל-1 לאורך 1.8 שנ': ערבוב בין גרסה מוצללת לצבע שטוח, דעיכת הפנים; תנועה מתמתנת (35%)"),
("L-3","0:04.7–0:05.6","כיתובים מופיעים: 'מתנ\"ס' (קפיצה), 'הר-ברכה' (עלייה), 'הר ברכה' (החלקה), 'עיר העתיד של השומרון' (עלייה)","e_back pop, fade+slide"),
("L-4","0:05.7–0:07.0","'חיבוק': שני הלוגואים מתקרבים ב-3.5% מרוחב המסך, ניצוצות צבעוניים וקו מפריד דק במרכז","פולסי ניצוצות במסלול סיבובי סביב נקודת המפגש"),
("L-5","0:08.3–0:09.0","פייד ללבן","הלוגואים נשארים על הדף יחד עד הסוף"),
]
for r in LOGO: ws.append(list(r))
ws.append([]);ws.append(["הערות","","",""])
for n in ["מקור לוגו מתנ\"ס: גרסת הדרייב (ללא לוגו המועצה האזורית שומרון). מקור לוגו הר ברכה: 2.png.","הפרדה לשכבות נעשית בקוד לפי צבע + רכיבים מחוברים (לא נדרש קובץ וקטורי)."]:
    ws.append(["•",n,"",""])
fmt(ws)
# ---------- 6 מוזיקה ופתוחים
ws=sheet(wb,"6 מוזיקה והחלטות",["נושא","המלצה / החלטה","סטטוס"],[24,100,20])
for r in [
("סוג הסרט","סיכום שנת הפעילות תשפ\"ו של המתנ\"ס (ההרשמה כבר נסגרה, השנה החדשה כבר החלה): ללא קריאה להרשמה. סיום: 'תודה על שנה מדהימה' · 'והשנה החדשה כבר כאן' · 'נתראה במתנ\"ס!'.","עודכן"),
("אורך","79 שנ' (בטווח 60–120), 39 סצנות, 43 תמונות ייחודיות + וידאו אחד.","עודכן"),
("כללי תמונות","כל תמונה מופיעה פעם אחת בלבד; אף תמונה לא יותר מ-3 שניות על המסך (המקסימום בפועל 3.0); הוסרו תמונות כפולות או דומות (מסומנות בקטלוג).","עודכן"),
("מוזיקה – סגנון","פופ-דאנס אופטימי ואנרגטי: קיק על כל פעימה, מחיאות, בס מתגלגל, סינת' סטאבים, הוק (מנגינה) בולט, סיידצ'יין, דרופים. פסנתר רגשי בחלק 'צמיחה'. ללא שירה. הרצועה הנוכחית מסונתזת – זמנית בלבד.","המלצה"),
("מוזיקה – BPM","120 BPM קבוע (פעימה = 0.5 שנ'; תיבה = 2 שנ'), בסול מז'ור (G–D–Em–C). כל החיתוכים והדרופים על פעימה. דרופ ראשון ב-0:16, דרופ סופי עם הלוגו ב-1:10.","המלצה"),
("מוזיקה – רישיון","צריך רצועה ברישיון לשימוש פומבי. מבקשים: pop/dance אופטימי, 120 BPM, עם דרופ אחרי ~16 שנ'.","נדרשת החלטה"),
("מספרים שיופיעו","600+ משתתפים בחוגים · 500 ילדים בקייטנות הקיץ · 70 תלמידי מוזיקה.","לאישור"),
("משחקייה","תיקיית 'משחקיה' ריקה – אין תמונות. אפשר להוסיף רגע אם יש תמונות.","חסר חומר"),
("פנים ילדים","זיהוי פנים ביומטרי לא בוצע; נדרש אישור הורים/המתנ\"ס לפרסום תמונות.","נדרש אישור"),
("לוגואים","שני הלוגואים מתחילים יחד בסצנה משותפת, ומתפצלים על המסך; הילדים של המתנ\"ס חיים (פנים, צל, תנועה) ובהמשך 'מתפשטים' לקווי הלוגו. לוגו המועצה האזורית שומרון הוסר.","בוצע"),
("שנה עברית","הטקסט הפותח 'שנת תשפ\"ו במתנ\"ס' – לאשר.","לאישור"),
]: ws.append(list(r))
fmt(ws)
wb.save("storyboard_har_bracha.xlsx"); print("saved",total)
